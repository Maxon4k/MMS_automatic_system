import os
import glob
from datetime import datetime

import openpyxl

# Діапазони (вкладка "Разом" у файлі небалансів)
CHECK_R1, CHECK_R2 = 4, 34      # рядки C4:Z34
CHECK_C1, CHECK_C2 = 3, 26      # колонки C..Z
SRC_C1 = 31                     # AE (джерело даних AE4:BB34)
GRAF_ROW_OFFSET = 12            # C4 -> C16 у файлі графіка (вкладка "ДД")

YELLOW_SUFFIX = "FFFF00"        # стандартний жовтий Excel


def _is_yellow(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return False
    rgb = f.start_color.rgb
    return isinstance(rgb, str) and rgb.upper().endswith(YELLOW_SUFFIX)


def find_todays_backup(backup_dir, suffix):
    """Шукає сьогоднішній бекап: 'Небаланси <Місяць> <Рік> MM DD <suffix>.xlsm'"""
    stamp = datetime.now().strftime("%m %d")
    pattern = os.path.join(backup_dir, f"*{stamp} {suffix}.xls*")
    files = glob.glob(pattern)
    if not files:
        return None
    return max(files, key=os.path.getmtime)


def check_and_update(backup_file, graf_file):
    """Якщо в C4:Z34 (Разом) є жовті комірки:
    AE4:BB34 -> C4:Z34 (у файлі бекапу) і ті самі значення -> C16:Z46 (ДД у графіку).
    Повертає True якщо були зміни."""
    if not os.path.exists(backup_file):
        print(f"❌ Бекап не знайдено: {backup_file}")
        return False
    if not os.path.exists(graf_file):
        print(f"❌ Файл графіка не знайдено: {graf_file}")
        return False

    keep_vba = backup_file.lower().endswith(".xlsm")

    # 1. Перевірка жовтих комірок
    wb_check = openpyxl.load_workbook(backup_file, keep_vba=keep_vba)
    if "Разом" not in wb_check.sheetnames:
        print("❌ Вкладку 'Разом' не знайдено в бекапі")
        wb_check.close()
        return False
    ws = wb_check["Разом"]

    yellow_cells = [
        (r, c)
        for r in range(CHECK_R1, CHECK_R2 + 1)
        for c in range(CHECK_C1, CHECK_C2 + 1)
        if _is_yellow(ws.cell(row=r, column=c))
    ]
    if not yellow_cells:
        print("ℹ️ Жовтих комірок немає — змін не потрібно.")
        wb_check.close()
        return False

    print(f"🟡 Знайдено жовтих комірок: {len(yellow_cells)}. Оновлюю дані.")

    # 2. Читаємо ЗНАЧЕННЯ AE4:BB34 окремим завантаженням (data_only),
    #    бо там можуть бути формули
    wb_vals = openpyxl.load_workbook(backup_file, data_only=True)
    ws_vals = wb_vals["Разом"]
    n_rows = CHECK_R2 - CHECK_R1 + 1
    n_cols = CHECK_C2 - CHECK_C1 + 1
    values = [
        [ws_vals.cell(row=CHECK_R1 + i, column=SRC_C1 + j).value for j in range(n_cols)]
        for i in range(n_rows)
    ]
    wb_vals.close()

    # 3. Вставляємо у C4:Z34 бекапу
    for i in range(n_rows):
        for j in range(n_cols):
            ws.cell(row=CHECK_R1 + i, column=CHECK_C1 + j).value = values[i][j]
    wb_check.save(backup_file)
    wb_check.close()
    print(f"✅ Оновлено C4:Z34 у {os.path.basename(backup_file)}")

    # 4. Ті самі значення -> графік, вкладка ДД, C16:Z46
    keep_vba_graf = graf_file.lower().endswith(".xlsm")
    wb_graf = openpyxl.load_workbook(graf_file, keep_vba=keep_vba_graf)
    if "ДД" not in wb_graf.sheetnames:
        print("❌ Вкладку 'ДД' не знайдено у графіку")
        wb_graf.close()
        return False
    ws_graf = wb_graf["ДД"]
    for i in range(n_rows):
        for j in range(n_cols):
            ws_graf.cell(row=CHECK_R1 + GRAF_ROW_OFFSET + i,
                         column=CHECK_C1 + j).value = values[i][j]
    wb_graf.save(graf_file)
    wb_graf.close()
    print(f"✅ Оновлено C16:Z46 у {os.path.basename(graf_file)}")
    return True


def run_for_company(cfg):
    """cfg: dict з ключами name, backup_dir, graf_path."""
    print(f"\n=== Перевірка змін: {cfg['name']} ===")
    backup = find_todays_backup(cfg["backup_dir"], cfg["name"])
    if not backup:
        print("❌ Сьогоднішній бекап не знайдено.")
        return False
    print(f"📄 Бекап: {os.path.basename(backup)}")
    return check_and_update(backup, cfg["graf_path"])