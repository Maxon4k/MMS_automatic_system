import os
import re
import shutil
import calendar
from datetime import date, datetime

import openpyxl
from openpyxl.styles import PatternFill

MONTHS_UA = {
    1: "Січень", 2: "Лютий", 3: "Березень", 4: "Квітень",
    5: "Травень", 6: "Червень", 7: "Липень", 8: "Серпень",
    9: "Вересень", 10: "Жовтень", 11: "Листопад", 12: "Грудень",
}

WEEKDAYS_UA = ["понеділок", "вівторок", "середа", "четвер", "п'ятниця", "субота", "неділя"]

# "Білий, фон 1, темніший 25%"
WEEKEND_FILL = PatternFill(fill_type="solid", start_color="FFBFBFBF", end_color="FFBFBFBF")
NO_FILL = PatternFill(fill_type=None)

# Блоки дат на "Разом":
# (стартовий рядок, кол. дня тижня, кол. дати, діапазон заливки: від-до включно)
RAZOM_DATE_BLOCKS = [
    (4,  1,  2,  (1, 27)),   # A4:AA34
    (39, 1,  2,  (1, 27)),   # A39:AA69
    (4,  29, 30, (29, 55)),  # AC4:BC34
    (39, 29, 30, (29, 55)),  # AC39:BC69
    (74, 1,  2,  (1, 27)),   # A74:AA104
]

SHEET_DATE_BLOCK = (40, 1, 2, (1, 27))  # A40:AA70 на інших вкладках
CLEAR_RANGE = (40, 70, 3, 26)           # C40:Z70 — очистка даних


def next_month(d: date):
    if d.month == 12:
        return d.year + 1, 1
    return d.year, d.month + 1


def fill_date_block(ws, start_row, col_wd, col_date, fill_span, year, month):
    """31 рядок: дати нового місяця + заливка вихідних по діапазону fill_span.
    Зайві рядки коротких місяців очищаються (значення і заливка)."""
    days_in_month = calendar.monthrange(year, month)[1]
    c_from, c_to = fill_span

    for i in range(31):
        row = start_row + i
        if i < days_in_month:
            d = date(year, month, i + 1)
            is_weekend = d.weekday() >= 5
            ws.cell(row=row, column=col_wd).value = WEEKDAYS_UA[d.weekday()]
            ws.cell(row=row, column=col_date).value = datetime(year, month, i + 1)
        else:
            is_weekend = False
            ws.cell(row=row, column=col_wd).value = None
            ws.cell(row=row, column=col_date).value = None

        fill = WEEKEND_FILL if is_weekend else NO_FILL
        for col in range(c_from, c_to + 1):
            ws.cell(row=row, column=col).fill = fill


def clear_data_range(ws, r1, r2, c1, c2):
    for row in range(r1, r2 + 1):
        for col in range(c1, c2 + 1):
            ws.cell(row=row, column=col).value = None


def update_title(ws, year, month):
    """Оновлює назву місяця (і рік) у заголовку B1 (об'єднаний B1:AA1)."""
    cell = ws.cell(row=1, column=2)
    val = cell.value
    if not isinstance(val, str):
        return
    new_val = val
    # заміна будь-якої назви місяця зі збереженням регістру першої літери
    for m_name in MONTHS_UA.values():
        def repl(match):
            word = MONTHS_UA[month]
            return word if match.group(0)[0].isupper() else word.lower()
        new_val = re.sub(m_name, repl, new_val, flags=re.IGNORECASE)
    # заміна року (напр. грудень 2026 -> січень 2027)
    new_val = re.sub(r"\b20\d{2}\b", str(year), new_val)
    if new_val != val:
        cell.value = new_val


def create_next_month_file(current_file_path):
    if not os.path.exists(current_file_path):
        print(f"❌ Файл не знайдено: {current_file_path}")
        return None

    today = date.today()
    ny, nm = next_month(today)

    folder = os.path.dirname(current_file_path)
    ext = os.path.splitext(current_file_path)[1]
    new_name = f"Небаланси {MONTHS_UA[nm]} {ny}{ext}"
    new_path = os.path.join(folder, new_name)

    if os.path.exists(new_path):
        print(f"ℹ️ Файл вже існує, пропускаю: {new_name}")
        return new_path

    print(f"📄 Створюю шаблон: {new_name}")
    shutil.copyfile(current_file_path, new_path)

    keep_vba = ext.lower() == ".xlsm"
    wb = openpyxl.load_workbook(new_path, keep_vba=keep_vba)

    # 1. Вкладка "Разом"
    if "Разом" in wb.sheetnames:
        ws = wb["Разом"]
        for start_row, col_wd, col_date, span in RAZOM_DATE_BLOCKS:
            fill_date_block(ws, start_row, col_wd, col_date, span, ny, nm)
        update_title(ws, ny, nm)
    else:
        print("⚠️ Вкладку 'Разом' не знайдено")

    # 2-3. Інші вкладки: дати + заливка A:AA, очистка C40:Z70, заголовок B1
    for sheet_name in wb.sheetnames:
        if sheet_name == "Разом":
            continue
        ws = wb[sheet_name]
        clear_data_range(ws, *CLEAR_RANGE)
        sr, cw, cd, span = SHEET_DATE_BLOCK
        fill_date_block(ws, sr, cw, cd, span, ny, nm)
        update_title(ws, ny, nm)

    wb.save(new_path)
    wb.close()
    print(f"✅ Шаблон готовий: {new_path}")
    return new_path


def is_last_day_of_month(d: date = None) -> bool:
    d = d or date.today()
    return d.day == calendar.monthrange(d.year, d.month)[1]


def rotate_export_folder(export_dir):
    """Перейменовує поточну папку export -> 'export <місяць що закінчився>'
    і створює нову порожню export."""
    today = date.today()
    month_name = MONTHS_UA[today.month].lower()
    parent = os.path.dirname(export_dir.rstrip("/\\"))
    archived = os.path.join(parent, f"export {month_name}")

    if os.path.exists(export_dir):
        if os.path.exists(archived):
            archived = os.path.join(parent, f"export {month_name} {today.year}")
        try:
            os.rename(export_dir, archived)
            print(f"📁 export -> {os.path.basename(archived)}")
        except PermissionError:
            print(f"⚠️ Не вдалось перейменувати export (папка зайнята). Пропускаю.")
            return
    os.makedirs(export_dir, exist_ok=True)
    print(f"📁 Створено нову папку: {export_dir}")


def run_monthly_maintenance(companies):
    """companies: список dict з ключами excel_path, download_dir, name.
    Створює шаблон наступного місяця + ротація export для кожної компанії."""
    results = []
    for c in companies:
        print(f"\n=== {c.get('name', '?')} ===")
        new_path = create_next_month_file(c['excel_path'])
        rotate_export_folder(c['download_dir'])
        results.append((c.get('name'), new_path))
    return results