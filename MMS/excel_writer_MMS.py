import openpyxl
import os
from datetime import datetime


def import_network_schedule(master_file_path, source_file_path):
    print(f"\n--- ІМПОРТ ГРАФІКУ ---")
    if not os.path.exists(source_file_path) or not os.path.exists(master_file_path):
        print("❌ Файли не знайдено")
        return

    wb_src = openpyxl.load_workbook(source_file_path, data_only=True)
    if "ДД" not in wb_src.sheetnames:
        wb_src.close()
        return

    ws_src = wb_src["ДД"]
    data = [[cell.value for cell in row] for row in ws_src["C16:Z46"]]
    wb_src.close()

    wb_master = openpyxl.load_workbook(master_file_path, keep_vba=True)
    ws_dest = wb_master["Разом"]
    for r, row in enumerate(data, start=4):
        for c, value in enumerate(row, start=3):
            ws_dest.cell(row=r, column=c, value=value)

    wb_master.save(master_file_path)
    wb_master.close()
    print("✅ Графік перенесено.")


def write_data_to_master(filepath, days_data):
    print(f"\n--- ЗАПИС ДАНИХ У EXCEL ---")
    try:
        wb = openpyxl.load_workbook(filepath, keep_vba=True)

        for target_date, data_dict, report, _, _ in days_data:
            # Шукаємо аркуш для зразка
            sample_sheet = next((s for s in wb.sheetnames if s in data_dict), wb.sheetnames[0])
            ws_ref = wb[sample_sheet]

            target_row = None
            for r in range(40, 80):
                cell_val = ws_ref.cell(row=r, column=2).value
                if isinstance(cell_val, datetime) and cell_val.date() == target_date.date():
                    target_row = r
                    break

            if not target_row:
                print(f"⚠️ Дата {target_date.date()} не знайдена в Excel. Пропускаємо цей день.")
                continue  # Йдемо до наступного дня, не перериваючи роботу

            # Записуємо дані
            for region, values in data_dict.items():
                if region in wb.sheetnames and report.get(region) != "MISSING_B":
                    ws = wb[region]
                    for i, val in enumerate(values):
                        ws.cell(row=target_row, column=3 + i).value = val

            print(f"✅ Записано дані за {target_date.strftime('%d.%m.%Y')}")

        # Зберігаємо файл ЛИШЕ ОДИН РАЗ після всіх циклів
        wb.save(filepath)
        wb.close()
        print(f"💾 Файл збережено: {os.path.basename(filepath)}")
    except Exception as e:
        print(f"❌ Помилка запису: {e}")